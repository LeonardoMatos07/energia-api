"""CSV and Excel files for the analyses."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.exporting.data import Analysis, Cell, Table, tables

FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")
UTF8_BOM = "﻿"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_SHEET_TITLE = 31


def safe_cell(value: Cell) -> Cell:
    """A text cell that starts like a formula would run as one when the file is opened, and some
    of these names come from public registries. Numbers are left alone."""
    if isinstance(value, str) and value.startswith(FORMULA_STARTS):
        return "'" + value
    return value


def to_csv(table: Table) -> bytes:
    """UTF-8 with a byte-order mark, which is what makes Excel read the accents right."""
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow(table.headers)
    for row in table.rows:
        writer.writerow(["" if cell is None else safe_cell(cell) for cell in row])
    return (UTF8_BOM + out.getvalue()).encode("utf-8")


def _fill_sheet(sheet: Worksheet, table: Table) -> None:
    sheet.append(list(table.headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E1B2C")
    for row in table.rows:
        sheet.append([safe_cell(cell) for cell in row])
    sheet.freeze_panes = "A2"
    for index, header in enumerate(table.headers, start=1):
        widest = max(
            [len(header)] + [len(str(row[index - 1])) for row in table.rows if row[index - 1]]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(widest + 2, 10), 60)


def to_xlsx(data: Analysis) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for table in tables(data).values():
        sheet = workbook.create_sheet(table.title[:MAX_SHEET_TITLE])
        _fill_sheet(sheet, table)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
